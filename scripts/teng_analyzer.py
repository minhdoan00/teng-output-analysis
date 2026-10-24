#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TENG Electrical Output Analysis & Visualization Suite (Voc, Isc, Isc -> Qsc, sigma)
Combines:
  1. Origin-style publication plotting & peak-to-peak statistical analysis (Voc, Isc, Qsc)
  2. Zero-Gap Midpoint Windowing short-circuit charge integration (Isc -> Qsc,max & sigma)
     with 4-panel diagnostic figures.
"""

import argparse
import datetime
import glob
import os
import re
import sys
import numpy as np
import pandas as pd
import matplotlib as mpl
import matplotlib.pyplot as plt
from matplotlib.ticker import MaxNLocator, AutoMinorLocator
from scipy.signal import find_peaks, peak_widths
from scipy.integrate import trapezoid, cumulative_trapezoid

# Ensure UTF-8 console output on Windows
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# ==============================================================================
# DEFAULT HARDWARE & ANALYSIS CONFIGURATION
# ==============================================================================
CSV_VOLTAGE_UNIT = "mV"       # Raw unit of Column 2: "V", "mV", "uV"
R_SHUNT_OHM = 100e3           # Shunt resistance: 100 kOhm
AMP_GAIN = 10.0               # Current preamplifier gain: x10
PROBE_ATTEN = 10.0            # Physical probe attenuation: x10
OSC_PROBE_SETTING = 10.0      # Oscilloscope probe setting: x10
AREA_CM2 = 25.0               # Active electrode area (cm^2)
VOC_MULTIPLIER = 1.0          # Voltage multiplier for Voc channel

DEFAULT_T_START = 1.0         # Default start time (s): discard first 1.0 s transient
DEFAULT_T_END = 7.0           # Default end time (s): keep next 6.0 s (1.0 -> 7.0 s)
T_WINDOW_MANUAL = (1.0, 7.0)  # Default (tmin, tmax); None = auto active-window
USE_ACTIVE_WINDOW = False     # True = adaptive active pulse window when window is None
MIN_PEAK_ISC_UA = 0.08        # Minimum peak height floor for Isc (uA) - adaptive 0.28*P98 handles higher signals
MIN_PEAK_VOC_V = None         # Minimum peak height floor for Voc (V)
MIN_DIST_S = 0.12             # Minimum distance between same-polarity peaks (s)
MAX_POS_WIDTH_S = None        # Optional maximum positive peak width filter for Voc (s)
EXCLUDE_EDGE_CYCLES = 1       # Edge cycles to drop in auto mode
EXCLUDE_EDGE_CYCLES_MANUAL = 0
MIN_REL_AMP = 0.5             # Drop cycles weaker than 50% of median amplitude

SHOW_PEAK_PEAK_TEXT = True    # Embed peak-to-peak text in top-left corner of Origin plot
PEAK_PEAK_COLOR = "#000000"   # Top-left text color for single-curve plot
X_FROM_ZERO = True            # Shift plotted time axis to start at 0 s
SHOW_PEAK_MARKERS = False     # Show peak markers on Origin plot
COLORS = ["#E00000", "#0000FF", "#008000", "#FF8C00", "#800080", "#000000", "#00A0A0", "#A52A2A"]
LINE_WIDTH = 2.2
FIG_SIZE = (6.6, 5.4)
DPI = 300

LABELS = {
    "en": {
        "t": "Time (s)",
        "Voc": "Voltage (V)",
        "Isc": "Current (\u00b5A)",
        "Qsc": "Transferred charge (nC)",
    },
    "vi": {
        "t": "Th\u1eddi gian (s)",
        "Voc": "\u0110i\u1ec7n \u00e1p (V)",
        "Isc": "D\u00f2ng \u0111i\u1ec7n (\u00b5A)",
        "Qsc": "\u0110i\u1ec7n t\u00edch (nC)",
    },
}
UNIT_V = {"V": 1.0, "mV": 1e-3, "uV": 1e-6}


# ==============================================================================
# 1. SCALING & DATA INGESTION
# ==============================================================================
def current_scale_uA(
    r_shunt_ohm=R_SHUNT_OHM,
    amp_gain=AMP_GAIN,
    probe_atten=PROBE_ATTEN,
    osc_probe_setting=OSC_PROBE_SETTING,
    csv_voltage_unit=CSV_VOLTAGE_UNIT,
):
    """Compute conversion scale from raw Column 2 reading to current in uA."""
    rtrans = r_shunt_ohm * amp_gain
    probe = osc_probe_setting / probe_atten
    return UNIT_V[csv_voltage_unit] * probe / rtrans * 1e6


def natural_sort_key(s):
    """Sort strings containing numbers in human/natural order (e.g., WaveData1239 before WaveData12310)."""
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", os.path.basename(str(s)))]


def load_oscilloscope_csv(path, return_meta=False):
    """Load oscilloscope CSV file, skipping '#' metadata header lines, and optionally extract #voltbase."""
    rows = []
    headers = []
    voltbase_uV = None
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f.read().replace("\r", "\n").split("\n"):
            s = line.strip()
            if not s:
                continue
            if "#" in s:
                headers.append(s)
                m = re.search(r"#voltbase\s*=\s*([0-9.]+)\s*\(\s*(uV|mV|V)\s*\)", s, re.I)
                if m:
                    val = float(m.group(1))
                    unit = m.group(2).lower()
                    voltbase_uV = val if unit == "uv" else (val * 1e3 if unit == "mv" else val * 1e6)
                else:
                    m2 = re.search(r"#voltbase\s*=\s*([0-9.]+)", s, re.I)
                    if m2:
                        voltbase_uV = float(m2.group(1))
                continue
            p = [x.strip() for x in s.split(",")]
            if len(p) < 2:
                continue
            try:
                rows.append((float(p[0]), float(p[1])))
            except ValueError:
                continue
    df = pd.DataFrame(rows, columns=["t_s", "raw"])
    if return_meta:
        return df, dict(voltbase_uV=voltbase_uV, headers=headers)
    return df


def auto_detect_signal_kind(path, df=None, meta=None, peer_vb_uV=None):
    """
    Classify an oscilloscope CSV file as 'Voc' (Voltage), 'Isc' (Current), or 'Qsc'
    with confidence assessment and physical evidence.

    Confidence levels:
      - 'EXPLICIT': explicitly stated in filename or file header comments.
      - 'HIGH': strong physical evidence (#voltbase >= 5V/div or <= 1V/div with matched amplitude).
      - 'MODERATE': inferred from paired channel context or amplitude without #voltbase.
      - 'AMBIGUOUS': intermediate scale (e.g. 2V/div) or borderline amplitude; user confirmation recommended.

    Returns:
      (kind, confidence, reason_str, evidence_dict)
    """
    bname = os.path.basename(str(path)).lower()
    evidence = {"filename": bname, "headers": []}

    if "qsc" in bname or "charge" in bname:
        return "Qsc", "EXPLICIT", "filename_keyword", evidence
    if bname.startswith("vol") or "_vol" in bname or "voc" in bname or "voltage" in bname:
        return "Voc", "EXPLICIT", "filename_keyword", evidence
    if bname.startswith("cur") or "_cur" in bname or "isc" in bname or "current" in bname:
        return "Isc", "EXPLICIT", "filename_keyword", evidence

    if df is None or meta is None:
        df, meta = load_oscilloscope_csv(path, return_meta=True)

    evidence["headers"] = meta.get("headers", [])
    for h in meta.get("headers", []):
        hl = h.lower()
        if "current" in hl or "isc" in hl:
            return "Isc", "EXPLICIT", "header_keyword", evidence
        if "voltage" in hl or "voc" in hl:
            return "Voc", "EXPLICIT", "header_keyword", evidence

    v = df["raw"].values if len(df) > 0 else np.array([0.0])
    if len(v) > 10:
        rs = pd.Series(v).rolling(5, center=True).std().fillna(999).values
        flat = v[rs <= max(np.percentile(rs, 55), 1e-12)]
        base = float(np.median(flat)) if len(flat) > 10 else float(np.median(v))
        ptp_raw = float(np.ptp(v - base))
    else:
        base, ptp_raw = 0.0, 0.0

    vb = meta.get("voltbase_uV")
    evidence["voltbase_uV"] = vb
    evidence["ptp_raw"] = ptp_raw
    evidence["baseline"] = base

    if vb is not None:
        vb_V = vb / 1e6
        evidence["voltbase_V_div"] = vb_V
        if vb >= 5_000_000:
            if ptp_raw >= 4000.0:
                return "Voc", "HIGH", f"voltbase={vb_V:g}V/div, ptp={ptp_raw:.0f}mV (Strong probability Voc)", evidence
            return "Voc", "MODERATE", f"voltbase={vb_V:g}V/div, lower ptp={ptp_raw:.0f}mV (Likely Voc, verify if weak sample)", evidence

        if vb <= 1_000_000:
            if ptp_raw <= 6000.0:
                return "Isc", "HIGH", f"voltbase={vb_V:g}V/div, ptp={ptp_raw:.0f}mV (Strong probability Isc via preamp)", evidence
            return "Isc", "MODERATE", f"voltbase={vb_V:g}V/div, large ptp={ptp_raw:.0f}mV (Likely Isc, verify preamp gain)", evidence

        # Intermediate 2 V/div scale
        if peer_vb_uV is not None:
            peer_V = peer_vb_uV / 1e6
            evidence["peer_voltbase_V_div"] = peer_V
            if peer_vb_uV >= 5_000_000:
                return "Isc", "HIGH", f"voltbase=2V/div paired with {peer_V:g}V/div channel (Likely Isc)", evidence
            if peer_vb_uV <= 500_000:
                return "Voc", "HIGH", f"voltbase=2V/div paired with {peer_V:g}V/div channel (Likely Voc)", evidence

        if ptp_raw >= 12000.0:
            return "Voc", "MODERATE", f"voltbase=2V/div, very large ptp={ptp_raw:.0f}mV (Probable Voc)", evidence
        if ptp_raw <= 4500.0:
            return "Isc", "MODERATE", f"voltbase=2V/div, small ptp={ptp_raw:.0f}mV (Probable Isc)", evidence
        return "Isc", "AMBIGUOUS", f"voltbase=2V/div, ptp={ptp_raw:.0f}mV without peer channel (Ambiguous: confirm with user)", evidence

    # Fallback when #voltbase is absent
    ptp_mV = ptp_raw * (UNIT_V.get(CSV_VOLTAGE_UNIT, 1e-3) / 1e-3)
    evidence["ptp_mV"] = ptp_mV
    if ptp_mV >= 12000.0:
        return "Voc", "MODERATE", f"No voltbase, large ptp={ptp_mV:.0f}mV (Probable Voc)", evidence
    if ptp_mV <= 3500.0:
        return "Isc", "MODERATE", f"No voltbase, small ptp={ptp_mV:.0f}mV (Probable Isc)", evidence
    return ("Voc" if ptp_mV >= 5500.0 else "Isc"), "AMBIGUOUS", f"No voltbase, intermediate ptp={ptp_mV:.0f}mV (Ambiguous: confirm with user)", evidence


def to_signal(df, kind):
    """Convert raw Column 2 reading to physical units (V for Voc, uA for Isc/Qsc)."""
    out = df.copy()
    scale = (
        UNIT_V[CSV_VOLTAGE_UNIT] * VOC_MULTIPLIER
        if kind == "Voc"
        else current_scale_uA()
    )
    out["y"] = out["raw"] * scale
    return out


# ==============================================================================
# 2. WINDOWING, RESTING-STATE DC BASELINE & PEAK PAIRING
# ==============================================================================
def active_window_and_baseline(df, window=None):
    """Slice time window and remove resting-state DC baseline offset."""
    t, v = df["t_s"].values, df["y"].values
    dt = float(np.mean(np.diff(t)))
    window = window if window is not None else T_WINDOW_MANUAL

    if window is not None:
        tmin, tmax = float(window[0]), float(window[1])
    elif USE_ACTIVE_WINDOW:
        win = max(5, int(round(0.05 / dt)))
        rstd = pd.Series(v).rolling(win, center=True).std().fillna(0).values
        rmed = pd.Series(v).rolling(win, center=True).median().bfill().ffill().values
        mid = v[int(0.15 * len(v)) : int(0.85 * len(v))]
        med = np.median(mid)
        iqr = max(np.percentile(np.abs(mid - med), 90), 1e-12)
        ok = (np.abs(rmed - med) < 25 * iqr) | (rstd < 1.5 * np.std(mid))
        idx = np.where(ok)[0]
        if len(idx) > 20:
            gaps = np.where(np.diff(idx) > int(round(0.5 / dt)))[0]
            seg = max(np.split(idx, gaps + 1), key=len)
            pad = int(round(0.02 / dt))
            tmin, tmax = t[max(0, seg[0] - pad)], t[min(len(t) - 1, seg[-1] + pad)]
        else:
            tmin, tmax = float(t[0]), float(t[-1])
    else:
        tmin = max(float(t[0]), DEFAULT_T_START)
        tmax = min(float(t[-1]), DEFAULT_T_END)

    sub = df[(df["t_s"] >= tmin) & (df["t_s"] <= tmax)].copy().reset_index(drop=True)
    if len(sub) < 20:
        raise ValueError(f"Time window [{tmin}, {tmax}] contains too few data points ({len(sub)}).")

    y = sub["y"].values
    rs = pd.Series(y).rolling(5, center=True).std().fillna(999).values
    flat = y[rs <= max(np.percentile(rs, 55), 1e-12)]
    base = float(np.median(flat)) if len(flat) > 10 else float(np.median(y))
    sub["yc"] = sub["y"] - base
    return sub, dict(tmin=tmin, tmax=tmax, baseline=base, dt=dt)


def detect_peaks(sub, dt, kind):
    """Detect positive and negative peaks using adaptive percentile thresholds."""
    v = sub["yc"].values
    dist = max(1, int(round(MIN_DIST_S / dt)))
    thr_p = 0.28 * np.percentile(np.maximum(v, 0), 98)
    thr_n = 0.28 * np.percentile(np.maximum(-v, 0), 98)
    floor = MIN_PEAK_ISC_UA if kind in ("Isc", "Qsc") else MIN_PEAK_VOC_V
    if floor is not None:
        thr_p, thr_n = max(thr_p, floor), max(thr_n, floor)
    pos, _ = find_peaks(v, height=thr_p, distance=dist)
    neg, _ = find_peaks(-v, height=thr_n, distance=dist)
    if kind == "Voc" and MAX_POS_WIDTH_S is not None and len(pos):
        w = peak_widths(v, pos, rel_height=0.5)[0] * dt
        pos = pos[w <= MAX_POS_WIDTH_S]
    return pos, neg


def pair_peaks(t, v, pos, neg):
    """Pair positive and negative peaks per cycle to compute peak-to-peak amplitude."""
    cols = ["cycle", "t_max_s", "max", "t_min_s", "min", "pp", "i_max", "i_min"]
    if len(pos) < 2 or len(neg) < 1:
        return pd.DataFrame(columns=cols)
    period = float(np.median(np.diff(t[pos])))
    rows = []
    for p in pos:
        cand = neg[np.abs(t[neg] - t[p]) <= 0.5 * period]
        if len(cand) == 0:
            continue
        q = cand[np.argmin(v[cand])]
        rows.append((0, t[p], v[p], t[q], v[q], v[p] - v[q], int(p), int(q)))
    df = pd.DataFrame(rows, columns=cols)
    df["cycle"] = np.arange(1, len(df) + 1)
    return df


# ==============================================================================
# 3. ZERO-GAP MIDPOINT WINDOWING & CHARGE INTEGRATION (Isc -> Qsc,max & sigma)
# ==============================================================================
def integrate_lobe(t, i, p, cs, ce, positive=True):
    """Integrate a single current lobe from peak p to zero-crossings within [cs, ce]."""
    y = i if positive else -i
    if p < cs or p > ce or y[p] <= 0:
        return 0.0, p, p, None
    l = p
    while l > cs and y[l - 1] > 0:
        l -= 1
    r = p
    while r < ce and y[r + 1] > 0:
        r += 1
    sp, ep = max(cs, l - 1), min(ce, r + 1)
    seg = np.maximum(y[sp : ep + 1], 0.0)
    tt = t[sp : ep + 1]
    return float(trapezoid(seg, tt)), sp, ep, cumulative_trapezoid(seg, tt, initial=0.0)


def charge_per_cycle(sub, pairs, area_cm2=AREA_CM2):
    """Integrate positive and negative charge lobes per paired cycle."""
    t = sub["t_s"].values
    i_nA = sub["yc"].values * 1000.0
    n = len(i_nA)
    first = np.minimum(pairs["i_max"].values, pairs["i_min"].values)
    last = np.maximum(pairs["i_max"].values, pairs["i_min"].values)
    mids = (last[:-1] + first[1:]) // 2 if len(pairs) > 1 else np.array([], dtype=int)
    half = int(np.median(first[1:] - last[:-1]) // 2) if len(pairs) > 1 else 0
    starts = np.r_[max(0, first[0] - half), mids + 1]
    ends = np.r_[mids, min(n - 1, last[-1] + half)]
    qp, qn = np.zeros(n), np.zeros(n)
    Qp, Qn = [], []
    for k in range(len(pairs)):
        cs, ce = int(starts[k]), int(ends[k])
        a, sp, ep, tr = integrate_lobe(t, i_nA, int(pairs.i_max[k]), cs, ce, True)
        b, sn, en, tn = integrate_lobe(t, i_nA, int(pairs.i_min[k]), cs, ce, False)
        if tr is not None:
            qp[sp : ep + 1] = np.maximum(qp[sp : ep + 1], tr)
        if tn is not None:
            qn[sn : en + 1] = np.maximum(qn[sn : en + 1], tn)
        Qp.append(a)
        Qn.append(b)
    pairs = pairs.copy()
    pairs["Q_pos_nC"], pairs["Q_neg_nC"] = Qp, Qn
    pairs["QSC_max_nC"] = np.maximum(Qp, Qn)
    pairs["sigma_uC_m2"] = pairs["QSC_max_nC"] * 1e-9 / (area_cm2 * 1e-4) * 1e6
    return pairs, np.maximum(qp, qn)


def _integrate_lobe_in_bounds(t, i_nA, peak_idx, c_start, c_end, is_positive=True):
    """Bounded zero-crossing lobe integration with 0.5% foot tolerance."""
    if peak_idx < c_start or peak_idx > c_end:
        return 0.0, peak_idx, peak_idx, np.array([0.0]), 0.0
    if is_positive:
        if i_nA[peak_idx] <= 0.0:
            return 0.0, peak_idx, peak_idx, np.array([0.0]), 0.0
        foot_tol = 0.005 * float(i_nA[peak_idx])
        left = peak_idx
        while left > c_start and i_nA[left - 1] > 0.0:
            if i_nA[left - 1] <= foot_tol and i_nA[left - 1] >= i_nA[left]:
                break
            left -= 1
        right = peak_idx
        while right < c_end and i_nA[right + 1] > 0.0:
            if i_nA[right + 1] <= foot_tol and i_nA[right + 1] >= i_nA[right]:
                break
            right += 1
        sp0, ep0 = max(c_start, left - 1), min(c_end, right + 1)
        seg_t = t[sp0 : ep0 + 1]
        seg_i = np.maximum(i_nA[sp0 : ep0 + 1], 0.0)
    else:
        if i_nA[peak_idx] >= 0.0:
            return 0.0, peak_idx, peak_idx, np.array([0.0]), 0.0
        foot_tol = 0.005 * float(-i_nA[peak_idx])
        left = peak_idx
        while left > c_start and i_nA[left - 1] < 0.0:
            if -i_nA[left - 1] <= foot_tol and i_nA[left - 1] <= i_nA[left]:
                break
            left -= 1
        right = peak_idx
        while right < c_end and i_nA[right + 1] < 0.0:
            if -i_nA[right + 1] <= foot_tol and i_nA[right + 1] <= i_nA[right]:
                break
            right += 1
        sp0, ep0 = max(c_start, left - 1), min(c_end, right + 1)
        seg_t = t[sp0 : ep0 + 1]
        seg_i = np.maximum(-i_nA[sp0 : ep0 + 1], 0.0)

    q_val = float(trapezoid(seg_i, seg_t))
    q_trace = cumulative_trapezoid(seg_i, seg_t, initial=0.0)
    dur_ms = float((t[ep0] - t[sp0]) * 1e3)
    return q_val, sp0, ep0, q_trace, dur_ms


def _best_lobe_in_window(t, i_nA, c_start, c_end, is_positive=True):
    """Find the primary positive or negative charge lobe within [c_start, c_end]."""
    win_slice = i_nA[c_start : c_end + 1]
    if len(win_slice) == 0:
        return 0.0, c_start, c_start, c_start, np.array([0.0]), 0.0
    p_ext_global = c_start + (int(np.argmax(win_slice)) if is_positive else int(np.argmin(win_slice)))
    best_q, best_sp0, best_ep0, best_trace, best_dur = _integrate_lobe_in_bounds(
        t, i_nA, p_ext_global, c_start, c_end, is_positive=is_positive
    )
    best_peak_global = p_ext_global
    mask = (win_slice > 0.0) if is_positive else (win_slice < 0.0)
    edges = np.diff(mask.astype(int))
    starts = np.where(edges == 1)[0] + 1
    ends = np.where(edges == -1)[0] + 1
    if mask[0]:
        starts = np.r_[0, starts]
    if mask[-1]:
        ends = np.r_[ends, len(win_slice) - 1]
    for s_rel, e_rel in zip(starts, ends):
        lobe_slice = win_slice[s_rel : e_rel + 1]
        if len(lobe_slice) == 0:
            continue
        p_rel = s_rel + (int(np.argmax(lobe_slice)) if is_positive else int(np.argmin(lobe_slice)))
        p_global = c_start + p_rel
        q_val, sp0, ep0, q_trace, dur_ms = _integrate_lobe_in_bounds(
            t, i_nA, p_global, c_start, c_end, is_positive=is_positive
        )
        if q_val > best_q:
            best_q, best_peak_global = q_val, p_global
            best_sp0, best_ep0 = sp0, ep0
            best_trace, best_dur = q_trace, dur_ms
    return best_q, best_peak_global, best_sp0, best_ep0, best_trace, best_dur


def analyze_teng_qsc_cycles(sub_df, area_cm2=AREA_CM2, min_peak_uA=MIN_PEAK_ISC_UA, min_distance_s=MIN_DIST_S):
    """
    Full v4 Zero-Gap Midpoint Windowing Qsc & sigma cycle analyzer:
      - Automatic phase-order detection ([- -> +] vs [+ -> -])
      - Zero-gap midpoint boundaries between consecutive cycles
      - Dual-polarity lobe integration (Q_+, |Q_-|, Q_SC,max, Sigma)
    """
    t = sub_df["t_s"].values
    i_uA = sub_df["yc"].values
    i_nA = i_uA * 1000.0
    n_samples = len(i_nA)
    dt = float(np.mean(np.diff(t)))
    min_dist_samples = max(1, int(round(min_distance_s / dt)))

    p98_pos = float(np.percentile(np.maximum(i_uA, 0.0), 98))
    p98_neg = float(np.percentile(np.maximum(-i_uA, 0.0), 98))
    thresh_pos = max(min_peak_uA, 0.28 * p98_pos)
    thresh_neg = max(min_peak_uA, 0.28 * p98_neg)

    peaks_pos, _ = find_peaks(i_uA, height=thresh_pos, distance=min_dist_samples)
    peaks_neg, _ = find_peaks(-i_uA, height=thresh_neg, distance=min_dist_samples)

    mean_pos = float(np.mean(i_uA[peaks_pos])) if len(peaks_pos) > 0 else 0.0
    mean_neg = float(np.mean(-i_uA[peaks_neg])) if len(peaks_neg) > 0 else 0.0

    if mean_neg > mean_pos and len(peaks_neg) >= 0.8 * len(peaks_pos):
        anchor_peaks, partner_peaks, anchor_is_pos = peaks_neg, peaks_pos, False
        dominant_polarity = "Negative (|Peak_neg| > Peak_pos)"
    else:
        anchor_peaks, partner_peaks, anchor_is_pos = peaks_pos, peaks_neg, True
        dominant_polarity = "Positive (Peak_pos >= |Peak_neg|)"

    if len(anchor_peaks) > 1:
        period_s = float(np.median(np.diff(t[anchor_peaks])))
    elif len(partner_peaks) > 1:
        period_s = float(np.median(np.diff(t[partner_peaks])))
    else:
        period_s = 0.22
    period_samples = max(4, int(round(period_s / dt)))

    dt_before_list, dt_after_list = [], []
    if len(anchor_peaks) > 0 and len(partner_peaks) > 0:
        for p_idx in anchor_peaks:
            prev_p = partner_peaks[partner_peaks < p_idx]
            next_p = partner_peaks[partner_peaks > p_idx]
            if len(prev_p) > 0:
                d_bef = float(t[p_idx] - t[prev_p[-1]])
                if d_bef < 0.85 * period_s:
                    dt_before_list.append(d_bef)
            if len(next_p) > 0:
                d_aft = float(t[next_p[0]] - t[p_idx])
                if d_aft < 0.85 * period_s:
                    dt_after_list.append(d_aft)

    med_before = float(np.median(dt_before_list)) if len(dt_before_list) > 0 else 0.5 * period_s
    med_after = float(np.median(dt_after_list)) if len(dt_after_list) > 0 else 0.5 * period_s
    partner_precedes_anchor = med_before < med_after
    pair_gap_samples = int(round((med_before if partner_precedes_anchor else med_after) / dt))
    pair_gap_samples = min(max(pair_gap_samples, int(0.10 * period_samples)), int(0.48 * period_samples))

    n_cycles = len(anchor_peaks)
    boundaries = np.zeros(n_cycles + 1, dtype=int)
    for k in range(n_cycles - 1):
        p_curr, p_next = int(anchor_peaks[k]), int(anchor_peaks[k + 1])
        local_period = p_next - p_curr
        local_gap = min(pair_gap_samples, int(0.48 * local_period))
        if partner_precedes_anchor:
            quiet_start, quiet_end = p_curr, max(p_curr + 1, p_next - local_gap)
        else:
            quiet_start, quiet_end = min(p_next - 1, p_curr + local_gap), p_next
        boundaries[k + 1] = int(round(0.5 * (quiet_start + quiet_end)))

    if n_cycles > 0:
        if partner_precedes_anchor:
            first_span = pair_gap_samples + int(round(0.5 * (period_samples - pair_gap_samples)))
            last_span = int(round(0.5 * (period_samples - pair_gap_samples)))
        else:
            first_span = int(round(0.5 * (period_samples - pair_gap_samples)))
            last_span = pair_gap_samples + int(round(0.5 * (period_samples - pair_gap_samples)))
        boundaries[0] = max(0, int(anchor_peaks[0]) - first_span)
        boundaries[n_cycles] = min(n_samples - 1, int(anchor_peaks[-1]) + last_span)

    if anchor_is_pos:
        phase_order_str = "[- -> +]" if partner_precedes_anchor else "[+ -> -]"
    else:
        phase_order_str = "[+ -> -]" if partner_precedes_anchor else "[- -> +]"

    area_m2 = area_cm2 * 1e-4
    cycles = []
    q_local_pos_nC = np.zeros_like(i_nA)
    q_local_neg_nC = np.zeros_like(i_nA)

    for idx in range(n_cycles):
        c_start, c_end = int(boundaries[idx]), int(boundaries[idx + 1])
        if c_end <= c_start:
            continue
        best_q_pos, p_pos_idx, sp0, ep0, q_trace_p, dur_pos_ms = _best_lobe_in_window(
            t, i_nA, c_start, c_end, is_positive=True
        )
        if best_q_pos > 0:
            q_local_pos_nC[sp0 : ep0 + 1] = np.maximum(q_local_pos_nC[sp0 : ep0 + 1], q_trace_p)

        best_q_neg, p_neg_idx, sn0, en0, q_trace_n, dur_neg_ms = _best_lobe_in_window(
            t, i_nA, c_start, c_end, is_positive=False
        )
        if best_q_neg > 0:
            q_local_neg_nC[sn0 : en0 + 1] = np.maximum(q_local_neg_nC[sn0 : en0 + 1], q_trace_n)

        p_max_win = c_start + int(np.argmax(i_uA[c_start : c_end + 1]))
        p_min_win = c_start + int(np.argmin(i_uA[c_start : c_end + 1]))
        if i_uA[p_max_win] > i_uA[p_pos_idx]:
            p_pos_idx = p_max_win
        if i_uA[p_min_win] < i_uA[p_neg_idx]:
            p_neg_idx = p_min_win

        i_max_uA = float(i_uA[p_pos_idx])
        i_min_uA = float(i_uA[p_neg_idx])
        q_sc_max_nC = max(best_q_pos, best_q_neg)
        q_sym_nC = 0.5 * (best_q_pos + best_q_neg)
        sigma_uCm2 = (q_sc_max_nC * 1e-9) / area_m2 * 1e6

        cycles.append({
            "Cycle": idx + 1,
            "t_start_s": round(float(t[c_start]), 4),
            "t_pos_s": round(float(t[p_pos_idx]), 4),
            "t_neg_s": round(float(t[p_neg_idx]), 4),
            "t_end_s": round(float(t[c_end]), 4),
            "I_pos_uA": round(i_max_uA, 3),
            "I_neg_uA": round(i_min_uA, 3),
            "I_pp_uA": round(i_max_uA - i_min_uA, 3),
            "Width_pos_ms": round(dur_pos_ms, 2),
            "Width_neg_ms": round(dur_neg_ms, 2),
            "Q_pos_nC": round(best_q_pos, 3),
            "Q_neg_nC": round(best_q_neg, 3),
            "Q_sym_nC": round(q_sym_nC, 3),
            "Q_SC_max_nC": round(q_sc_max_nC, 3),
            "Sigma_uCm2": round(sigma_uCm2, 3),
        })

    cycles_df = pd.DataFrame(cycles)
    if len(cycles_df) >= 4:
        mid_slice = cycles_df.iloc[1:-1]
        mid_qmax_med = float(mid_slice["Q_SC_max_nC"].median())
        status = []
        for i, row in cycles_df.iterrows():
            is_edge = i == 0 or i == len(cycles_df) - 1
            if is_edge and row["Q_SC_max_nC"] < 0.75 * mid_qmax_med:
                status.append("Edge")
            else:
                status.append("Steady")
        cycles_df["Status"] = status
    elif len(cycles_df) > 0:
        cycles_df["Status"] = "Steady"

    sub_out = sub_df.copy()
    sub_out["Q_pos_trace_nC"] = q_local_pos_nC
    sub_out["Q_neg_trace_nC"] = q_local_neg_nC
    sub_out["Q_SC_trace_nC"] = np.maximum(q_local_pos_nC, q_local_neg_nC)
    return cycles_df, sub_out, f"{dominant_polarity} | Phase: {phase_order_str}"


# ==============================================================================
# 4. PIPELINE & STATISTICAL AGGREGATION
# ==============================================================================
def analyze_file(path, kind="auto", window=None, preview=False, peer_vb_uV=None):
    """Analyze a single oscilloscope CSV file for Voc, Isc, or Qsc (with probabilistic confidence assessment)."""
    df_raw, meta = load_oscilloscope_csv(path, return_meta=True)
    if kind is None or str(kind).lower() == "auto":
        kind, confidence, detect_reason, evidence = auto_detect_signal_kind(
            path, df=df_raw, meta=meta, peer_vb_uV=peer_vb_uV
        )
    else:
        confidence = "EXPLICIT"
        detect_reason = "explicit"
        evidence = {"explicit_kind": kind}
    df = to_signal(df_raw, kind)
    sub, info = active_window_and_baseline(df, window)
    info["confidence"] = confidence
    info["detect_reason"] = detect_reason
    info["evidence"] = evidence
    info["voltbase_uV"] = meta.get("voltbase_uV")
    pos, neg = detect_peaks(sub, info["dt"], kind)
    t, v = sub["t_s"].values, sub["yc"].values
    pairs = pair_peaks(t, v, pos, neg)
    qtrace = None
    if kind == "Qsc" and len(pairs):
        pairs, qtrace = charge_per_cycle(sub, pairs)
    key = "QSC_max_nC" if kind == "Qsc" else "pp"
    n_all = len(pairs)
    manual = ((window is not None) or (T_WINDOW_MANUAL is not None)) and not preview
    n_edge = EXCLUDE_EDGE_CYCLES_MANUAL if manual else EXCLUDE_EDGE_CYCLES
    if n_edge > 0 and n_all > 2 * n_edge + 1:
        pairs = pairs.iloc[n_edge : n_all - n_edge].reset_index(drop=True)
    if MIN_REL_AMP is not None and len(pairs):
        pairs = pairs[pairs[key] >= MIN_REL_AMP * pairs[key].median()].reset_index(drop=True)
    pairs["cycle"] = np.arange(1, len(pairs) + 1)
    w = window if window is not None else T_WINDOW_MANUAL
    x0 = 0.0 if preview else ((w[0] if w is not None else sub["t_s"].values[0]) if X_FROM_ZERO else 0.0)
    x = sub["t_s"].values - x0
    y = qtrace if kind == "Qsc" and qtrace is not None else v
    xlim = (w[0], w[1]) if preview else ((0.0, w[1] - w[0]) if (w is not None and X_FROM_ZERO) else None)
    return dict(
        path=path,
        kind=kind,
        confidence=confidence,
        detect_reason=detect_reason,
        evidence=evidence,
        x=x,
        y=y,
        full_df=df,
        sub=sub,
        info=info,
        pairs=pairs,
        n_pos=len(pos),
        n_neg=len(neg),
        n_removed=n_all - len(pairs),
        x0=x0,
        xlim=xlim,
    )


def stats_of(res, kind, name):
    """Compute Mean and SD summary rows for a processed result dict."""
    P = res["pairs"]

    def ms(c):
        return (
            (float(P[c].mean()), float(P[c].std(ddof=1)) if len(P) > 1 else np.nan)
            if len(P)
            else (np.nan, np.nan)
        )

    if kind == "Qsc":
        rows = [
            ("Mean Q+ (nC)", *ms("Q_pos_nC")),
            ("Mean Q- (nC)", *ms("Q_neg_nC")),
            ("Mean QSC,max (nC)", *ms("QSC_max_nC")),
            ("Mean sigma (\u00b5C/m2)", *ms("sigma_uC_m2")),
        ]
    else:
        u = "V" if kind == "Voc" else "\u00b5A"
        rows = [
            (f"Mean peak max ({u})", *ms("max")),
            (f"Mean peak min ({u})", *ms("min")),
            (f"Mean peak-peak ({u})", *ms("pp")),
        ]
    return [
        dict(Data=name, Parameter=r[0], Mean=r[1], Std=r[2], N_cycles=len(P))
        for r in rows
    ]


# ==============================================================================
# 5. VISUALIZATION ENGINES (ORIGIN-STYLE & 4-PANEL DIAGNOSTIC)
# ==============================================================================
def origin_plot(results, names, kind, lang, save_path, show_pp=None):
    """Generate a publication-grade Origin-style plot with top-left peak-to-peak annotation."""
    if show_pp is None:
        show_pp = SHOW_PEAK_PEAK_TEXT
    lbl_dict = LABELS.get(lang, LABELS["en"])

    mpl.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial", "Liberation Sans", "DejaVu Sans"],
        "font.weight": "bold",
        "axes.labelweight": "bold",
        "axes.linewidth": 2.0,
    })
    fig, ax = plt.subplots(figsize=FIG_SIZE, dpi=150)
    n = len(results)
    ymin, ymax, xmax = np.inf, -np.inf, 0
    for k, (res, name) in enumerate(zip(results, names)):
        c = COLORS[k % len(COLORS)]
        ax.plot(res["x"], res["y"], color=c, lw=LINE_WIDTH, label=name)
        ymin = min(ymin, float(np.min(res["y"])))
        ymax = max(ymax, float(np.max(res["y"])))
        xmax = max(xmax, float(res["x"][-1]))
        if SHOW_PEAK_MARKERS and kind != "Qsc":
            P = res["pairs"]
            x0 = res["x0"]
            ax.plot(P["t_max_s"] - x0, P["max"], "o", mfc="none", mec=c, ms=6)
            ax.plot(P["t_min_s"] - x0, P["min"], "s", mfc="none", mec=c, ms=6)

    ax.set_xlabel(lbl_dict["t"], fontsize=24, fontweight="bold", labelpad=6)
    ax.set_ylabel(lbl_dict[kind], fontsize=24, fontweight="bold", labelpad=6)
    ax.xaxis.set_major_locator(MaxNLocator(nbins=5, steps=[1, 2, 5, 10]))
    ax.yaxis.set_major_locator(MaxNLocator(nbins=5, steps=[1, 2, 5, 10]))
    ax.xaxis.set_minor_locator(AutoMinorLocator(2))
    ax.yaxis.set_minor_locator(AutoMinorLocator(2))
    ax.tick_params(which="major", direction="in", length=8, width=2, labelsize=20)
    ax.tick_params(which="minor", direction="in", length=4, width=1.5)
    for lb in ax.get_xticklabels() + ax.get_yticklabels():
        lb.set_fontweight("bold")

    if results[0]["xlim"] is not None:
        ax.set_xlim(*results[0]["xlim"])
    else:
        ax.set_xlim(0 if X_FROM_ZERO else min(r["x"][0] for r in results), xmax)

    head = 0.18 + 0.09 * n if n > 1 else 0.22
    span = max(ymax - ymin, 1e-9)
    ax.set_ylim(ymin - 0.08 * span, ymax + head * span)

    # 1) Top-Right: Sample Name or Legend
    if n == 1:
        ax.text(
            0.97, 0.97, names[0], transform=ax.transAxes,
            ha="right", va="top", fontsize=18, fontweight="bold"
        )
    else:
        leg = ax.legend(
            loc="upper right", frameon=False, handlelength=1.6,
            prop={"weight": "bold", "size": 17}
        )
        for txt, ln in zip(leg.get_texts(), leg.get_lines()):
            txt.set_color(ln.get_color())

    # 2) Top-Left: Peak-to-Peak Annotation
    if show_pp:
        y_top = 0.97
        dy = 0.075 if n > 1 else 0.0
        for k, (res, nm) in enumerate(zip(results, names)):
            P = res.get("pairs")
            pp_col = "QSC_max_nC" if kind == "Qsc" else "pp"
            unit = "V" if kind == "Voc" else ("\u00b5A" if kind == "Isc" else "nC")
            sym = "V_{pp}" if kind == "Voc" else ("I_{pp}" if kind == "Isc" else "Q_{sc}")
            if P is not None and len(P) > 0:
                m_val = float(P[pp_col].mean())
                s_val = float(P[pp_col].std(ddof=1)) if len(P) > 1 else np.nan
                if not np.isnan(s_val) and s_val > 0:
                    pp_txt = f"${sym}$ = {m_val:.2f} \u00b1 {s_val:.2f} {unit}"
                else:
                    pp_txt = f"${sym}$ = {m_val:.2f} {unit}"
            else:
                raw_ptp = float(np.ptp(res["y"]))
                pp_txt = f"${sym}$ ~ {raw_ptp:.2f} {unit}"

            txt_color = COLORS[k % len(COLORS)] if n > 1 else (PEAK_PEAK_COLOR or COLORS[0])
            ax.text(
                0.04, y_top - k * dy, pp_txt, transform=ax.transAxes,
                ha="left", va="top", fontsize=18 if n <= 2 else 15,
                fontweight="bold", color=txt_color
            )

    fig.tight_layout()
    fig.savefig(save_path, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    return save_path


def plot_qsc_4panel(full_df, sub_df_q, cycles_df, info, filename_label, save_path, lang="vi"):
    """
    Generate the 4-panel (300 DPI) diagnostic figure for Isc -> Qsc,max & sigma integration:
      (a) Full raw oscilloscope trace & active window + DC baseline
      (b) Baseline-corrected I_SC(t) with positive & negative peak markers
      (c) Zoomed view of 3 central cycles with shaded Q_+, |Q_-| areas & cumulative charge curves
      (d) Cycle-by-cycle bar chart of Q_+, |Q_-|, and Q_SC,max trajectory
    """
    vi = lang == "vi"
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.weight": "normal",
        "axes.labelweight": "normal",
        "axes.linewidth": 1.0,
        "font.size": 9.5,
        "axes.titlesize": 10.5,
        "axes.labelsize": 9.5,
        "figure.dpi": 150,
        "savefig.dpi": 300,
    })

    steady_df = cycles_df[cycles_df["Status"] == "Steady"] if "Status" in cycles_df else cycles_df
    if len(steady_df) == 0:
        steady_df = cycles_df

    q_mean = float(steady_df["Q_SC_max_nC"].mean()) if len(steady_df) else 0.0
    q_std = float(steady_df["Q_SC_max_nC"].std()) if len(steady_df) > 1 else 0.0
    q_max = float(steady_df["Q_SC_max_nC"].max()) if len(steady_df) else 0.0
    sigma_mean = float(steady_df["Sigma_uCm2"].mean()) if len(steady_df) else 0.0
    i_pos_max = float(steady_df["I_pos_uA"].max()) if len(steady_df) else 0.0
    i_pos_mean = float(steady_df["I_pos_uA"].mean()) if len(steady_df) else 0.0
    i_neg_min = float(steady_df["I_neg_uA"].min()) if len(steady_df) else 0.0
    i_neg_mean = float(steady_df["I_neg_uA"].mean()) if len(steady_df) else 0.0
    base_uA = info["baseline"]

    fig, axes = plt.subplots(2, 2, figsize=(14.4, 9.4))
    title_str = (
        f"K\u1ebeT QU\u1ea2 PH\u00c2N T\u00cdCH D\u00d2NG NG\u1eaeN M\u1ea0CH $I_{{SC}}(t)$ & \u0110I\u1ec6N T\u00cdCH $Q_{{SC,max}}$ \u2014 {filename_label}"
        if vi else
        f"SHORT-CIRCUIT CURRENT $I_{{SC}}(t)$ & TRANSFERRED CHARGE $Q_{{SC,max}}$ ANALYSIS \u2014 {filename_label}"
    )
    fig.suptitle(title_str, fontsize=12.5, fontweight="bold", y=0.98)

    # Panel (a)
    ax1 = axes[0, 0]
    ax1.plot(full_df["t_s"], full_df["y"], color="#4C566A", lw=0.8,
             label="T\u00edn hi\u1ec7u g\u1ed1c $I_{raw}(t)$" if vi else "Raw signal $I_{raw}(t)$")
    ax1.axhline(base_uA, color="#D62728", ls="--", lw=1.3,
                label=f"N\u1ec1n DC: {base_uA:.3f} $\\mu$A" if vi else f"DC baseline: {base_uA:.3f} $\\mu$A")
    ax1.axvspan(info["tmin"], info["tmax"], color="#A3BE8C", alpha=0.20,
                label=f"C\u1eeda s\u1ed5 ph\u00e2n t\u00edch ({info['tmin']:.2f} - {info['tmax']:.2f} s)" if vi
                else f"Analysis window ({info['tmin']:.2f} - {info['tmax']:.2f} s)")
    ax1.set_title("(a) To\u00e0n b\u1ed9 b\u1ea3n ghi & C\u1eeda s\u1ed5 ph\u00e2n t\u00edch" if vi else "(a) Full Recording & Active Window", fontweight="bold")
    ax1.set_xlabel("Th\u1eddi gian $t$ (s)" if vi else "Time $t$ (s)")
    ax1.set_ylabel("$I_{raw}$ ($\\mu$A)")
    ax1.grid(True, ls=":", alpha=0.6)
    ax1.legend(loc="lower left", fontsize=8.2)

    # Panel (b)
    ax2 = axes[0, 1]
    ax2.plot(sub_df_q["t_s"], sub_df_q["yc"], color="#1F77B4", lw=0.85, label="$I_{SC}(t) = I_{raw}(t) - I_{baseline}$")
    if len(cycles_df):
        ax2.scatter(cycles_df["t_pos_s"], cycles_df["I_pos_uA"], color="#2CA02C", s=22, zorder=5,
                    label=f"Peak + (Mean: +{i_pos_mean:.2f} $\\mu$A | Max: +{i_pos_max:.2f} $\\mu$A)")
        ax2.scatter(cycles_df["t_neg_s"], cycles_df["I_neg_uA"], color="#D62728", marker="v", s=24, zorder=5,
                    label=f"Peak - (Mean: {i_neg_mean:.2f} $\\mu$A | Min: {i_neg_min:.2f} $\\mu$A)")
    ax2.axhline(0, color="black", lw=0.8)
    ax2.set_title(f"(b) D\u00f2ng ng\u1eafn m\u1ea1ch $I_{{SC}}(t)$ sau kh\u1eed n\u1ec1n ({len(cycles_df)} chu k\u1ef3)" if vi
                  else f"(b) Baseline-Corrected $I_{{SC}}(t)$ ({len(cycles_df)} cycles)", fontweight="bold")
    ax2.set_xlabel("Th\u1eddi gian $t$ (s)" if vi else "Time $t$ (s)")
    ax2.set_ylabel("$I_{SC}$ ($\\mu$A)")
    ax2.grid(True, ls=":", alpha=0.6)
    ax2.legend(loc="upper right", fontsize=8.0)

    # Panel (c)
    ax3 = axes[1, 0]
    if len(cycles_df) > 0:
        mid_idx = len(cycles_df) // 2
        c_from = max(0, mid_idx - 1)
        c_to = min(len(cycles_df) - 1, c_from + 2)
        t_z1 = float(cycles_df.loc[c_from, "t_start_s"])
        t_z2 = float(cycles_df.loc[c_to, "t_end_s"])
        zm = (sub_df_q["t_s"] >= t_z1) & (sub_df_q["t_s"] <= t_z2)
        tz = sub_df_q.loc[zm, "t_s"].values
        iz = sub_df_q.loc[zm, "yc"].values
        qz_pos = sub_df_q.loc[zm, "Q_pos_trace_nC"].values
        qz_neg = sub_df_q.loc[zm, "Q_neg_trace_nC"].values

        ax3_t = ax3.twinx()
        ax3.plot(tz, iz, color="#1F77B4", lw=1.35, label="$I_{SC}(t)$ ($\\mu$A)")
        ax3.fill_between(tz, 0, np.maximum(iz, 0), where=(iz > 0), color="#2CA02C", alpha=0.25, label="$Q_+$ area")
        ax3.fill_between(tz, np.minimum(iz, 0), 0, where=(iz < 0), color="#FF7F0E", alpha=0.25, label="$|Q_-|$ area")
        ax3_t.plot(tz, qz_pos, color="#D62728", lw=1.6, ls="-.", label="$Q_+(t)$ (nC)")
        ax3_t.plot(tz, qz_neg, color="#9467BD", lw=1.6, ls="--", label="$|Q_-(t)|$ (nC)")
        for ck_i in range(c_from, c_to + 1):
            ax3.axvline(float(cycles_df.loc[ck_i, "t_start_s"]), color="gray", ls=":", lw=0.8, alpha=0.6)
        ax3.axvline(float(cycles_df.loc[c_to, "t_end_s"]), color="gray", ls=":", lw=0.8, alpha=0.6)
        ax3_t.set_ylabel("Transferred Charge $Q(t)$ (nC)", color="#D62728")
        l1, lb1 = ax3.get_legend_handles_labels()
        l2, lb2 = ax3_t.get_legend_handles_labels()
        ax3.legend(l1 + l2, lb1 + lb2, loc="upper left", fontsize=7.8)
        ax3.set_title(f"(c) Ph\u00f3ng \u0111\u1ea1i Chu k\u1ef3 {c_from+1} - {c_to+1} & T\u00edch ph\u00e2n $Q_+(t), |Q_-(t)|$" if vi
                      else f"(c) Zoomed Cycles {c_from+1} - {c_to+1} & Cumulative Charge $Q_+(t), |Q_-(t)|$", fontweight="bold")
    ax3.axhline(0, color="black", lw=0.7)
    ax3.set_xlabel("Th\u1eddi gian $t$ (s)" if vi else "Time $t$ (s)")
    ax3.set_ylabel("$I_{SC}$ ($\\mu$A)", color="#1F77B4")
    ax3.grid(True, ls=":", alpha=0.6)

    # Panel (d)
    ax4 = axes[1, 1]
    if len(cycles_df) > 0:
        ck = cycles_df["Cycle"].values
        w = 0.36
        ax4.bar(ck - w / 2, cycles_df["Q_pos_nC"], width=w, color="#2CA02C", alpha=0.85, label="$Q_+$ (nC)")
        ax4.bar(ck + w / 2, cycles_df["Q_neg_nC"], width=w, color="#FF7F0E", alpha=0.85, label="$|Q_-|$ (nC)")
        ax4.plot(ck, cycles_df["Q_SC_max_nC"], "o-", color="#1F77B4", lw=1.3, ms=4.0, label="$Q_{SC,max} = \\max(Q_+, |Q_-|)$")
        ax4.axhline(q_mean, color="#D62728", ls="--", lw=1.6,
                    label=f"Mean: {q_mean:.2f} $\\pm$ {q_std:.2f} nC ($\\sigma = {sigma_mean:.2f}\\ \\mu$C/m$^2$)")
        ax4.set_ylim(0, max(10.0, q_max * 1.28))
        ax4.legend(loc="upper right", fontsize=8.0)
    ax4.set_title(f"(d) \u0110i\u1ec7n t\u00edch $Q_{{SC,max}}$ qua c\u00e1c chu k\u1ef3 (Max: {q_max:.2f} nC)" if vi
                  else f"(d) Cycle-by-Cycle $Q_{{SC,max}}$ Distribution (Max: {q_max:.2f} nC)", fontweight="bold")
    ax4.set_xlabel("Chu k\u1ef3 (Cycle #)" if vi else "Cycle #")
    ax4.set_ylabel("$Q_{SC}$ (nC)")
    ax4.grid(True, axis="y", ls=":", alpha=0.6)

    plt.tight_layout(rect=[0, 0, 1, 0.95])
    plt.savefig(save_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return save_path


# ==============================================================================
# 6. BATCH PROCESSING (FOLDER OF CSVs & MULTI-COLUMN EXCEL WORKBOOK)
# ==============================================================================
def batch_process_folder(
    folder_path="Data/Vol-Cur",
    out_dir=None,
    window=(DEFAULT_T_START, DEFAULT_T_END),
    lang="vi",
    date_tag=None,
):
    """Batch-process all oscilloscope CSV files in a folder with automatic Voc/Isc recognition."""
    if date_tag is None:
        date_tag = datetime.date.today().strftime("%Y%m%d")
    if out_dir is None:
        out_dir = os.path.join(folder_path, "draft_png")
    os.makedirs(out_dir, exist_ok=True)

    csv_files = sorted(
        [
            f for f in os.listdir(folder_path)
            if f.lower().endswith(".csv")
            and not f.startswith("202")
            and "summary" not in f.lower()
            and "dataset" not in f.lower()
            and not f.lower().endswith("_origin_xy.csv")
            and not f.lower().endswith("_cycles.csv")
        ],
        key=natural_sort_key,
    )

    # Pre-load raw data & metadata to support paired channel context (#voltbase comparison)
    loaded = []
    for f in csv_files:
        full_path = os.path.join(folder_path, f)
        df_raw, meta = load_oscilloscope_csv(full_path, return_meta=True)
        if len(df_raw) >= 20:
            loaded.append((f, full_path, df_raw, meta))

    summary_rows = []
    for i, (f, full_path, df_raw, meta) in enumerate(loaded):
        base_name = os.path.splitext(f)[0]
        peer_idx = i + 1 if i % 2 == 0 else i - 1
        peer_vb = loaded[peer_idx][3].get("voltbase_uV") if 0 <= peer_idx < len(loaded) else None

        kind, confidence, detect_reason, evidence = auto_detect_signal_kind(
            full_path, df=df_raw, meta=meta, peer_vb_uV=peer_vb
        )
        res = analyze_file(full_path, kind=kind, window=window, preview=False, peer_vb_uV=peer_vb)
        res["confidence"] = confidence
        res["detect_reason"] = detect_reason

        low_base = base_name.lower()
        if low_base.startswith("vol") or low_base.startswith("cur") or low_base.startswith("qsc"):
            tag_name = base_name
        else:
            prefix = "Vol" if kind == "Voc" else ("Cur" if kind == "Isc" else "Qsc")
            tag_name = f"{prefix}_{base_name}"

        out_png_name = f"{date_tag}_uet_draft_plot_{tag_name}_v1.png"
        out_png_path = os.path.join(out_dir, out_png_name)
        origin_plot([res], [base_name], kind, lang, out_png_path, show_pp=True)

        P = res["pairs"]
        pp_col = "QSC_max_nC" if kind == "Qsc" else "pp"
        unit = "V" if kind == "Voc" else ("\u00b5A" if kind == "Isc" else "nC")
        if len(P) > 0:
            m_pp = float(P[pp_col].mean())
            s_pp = float(P[pp_col].std(ddof=1)) if len(P) > 1 else np.nan
            m_pos = float(P["max"].mean()) if "max" in P else np.nan
            m_neg = float(P["min"].mean()) if "min" in P else np.nan
            n_cyc = len(P)
        else:
            m_pp = float(np.ptp(res["y"]))
            s_pp = np.nan
            m_pos = float(np.max(res["y"]))
            m_neg = float(np.min(res["y"]))
            n_cyc = 0

        summary_rows.append({
            "Data": base_name,
            "Type": kind,
            "Confidence": confidence,
            "Detection_Rule": detect_reason,
            "Peak_Peak_Mean": round(m_pp, 3),
            "Peak_Peak_Std": round(s_pp, 3) if not np.isnan(s_pp) else np.nan,
            "Unit": unit,
            "N_cycles": n_cyc,
            "Peak_Pos_Mean": round(m_pos, 3),
            "Peak_Neg_Mean": round(m_neg, 3),
            "Baseline_DC": round(res["info"]["baseline"], 4),
            "Window_Start_s": round(res["info"]["tmin"], 3),
            "Window_End_s": round(res["info"]["tmax"], 3),
            "Image_File": out_png_name,
        })

    sum_df = pd.DataFrame(summary_rows)
    sum_csv_path = os.path.join(
        folder_path, f"{date_tag}_uet_dataset_vol_cur_peak_peak_{len(summary_rows)}files_v1.csv"
    )
    sum_df.to_csv(sum_csv_path, index=False, encoding="utf-8-sig")
    return sum_df


def batch_process_excel(
    excel_path="Data/Book1.xlsx",
    out_dir="Data/NIPS",
    window=(DEFAULT_T_START, DEFAULT_T_END),
    lang="vi",
    date_tag=None,
):
    """Batch-process a multi-column Excel workbook into Origin plots, CSVs, and summary table."""
    import openpyxl

    if date_tag is None:
        date_tag = datetime.date.today().strftime("%Y%m%d")
    draft_png_dir = os.path.join(out_dir, "draft_png")
    csv_dir = os.path.join(out_dir, "csv")
    os.makedirs(draft_png_dir, exist_ok=True)
    os.makedirs(csv_dir, exist_ok=True)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    summary_rows = []

    def _clean(s):
        r = str(s).replace("@", "-").replace("*", "-star").replace("/", "").replace(" ", "-")
        while "--" in r:
            r = r.replace("--", "-")
        return r.strip("-")

    for c in range(1, ws.max_column, 2):
        y_header = ws.cell(1, c + 1).value
        y_unit = ws.cell(2, c + 1).value
        raw_name = ws.cell(3, c + 1).value
        if not raw_name or not y_header:
            continue
        raw_name = str(raw_name).strip()
        kind = "Voc" if "volt" in str(y_header).lower() or str(y_unit).strip() == "V" else "Isc"
        prefix = "Vol" if kind == "Voc" else "Cur"
        base_file_name = f"{prefix}_{_clean(raw_name)}"

        times, raws = [], []
        for r in range(4, ws.max_row + 1):
            tv, yv = ws.cell(r, c).value, ws.cell(r, c + 1).value
            if tv is not None and yv is not None:
                times.append(float(tv))
                raws.append(float(yv))
        if not times:
            continue

        df_raw = pd.DataFrame({"t_s": times, "raw": raws})
        df_raw.to_csv(os.path.join(csv_dir, f"{base_file_name}.csv"), index=False)

        df_sig = to_signal(df_raw, kind)
        sub, info = active_window_and_baseline(df_sig, window=window)
        pos, neg = detect_peaks(sub, info["dt"], kind)
        t, v = sub["t_s"].values, sub["yc"].values
        pairs = pair_peaks(t, v, pos, neg)
        if MIN_REL_AMP is not None and len(pairs):
            pairs = pairs[pairs["pp"] >= MIN_REL_AMP * pairs["pp"].median()].reset_index(drop=True)
        pairs["cycle"] = np.arange(1, len(pairs) + 1)
        x0 = window[0] if (window is not None and X_FROM_ZERO) else 0.0
        x = sub["t_s"].values - x0
        xlim = (0.0, window[1] - window[0]) if (window is not None and X_FROM_ZERO) else window
        res = dict(x=x, y=v, sub=sub, info=info, pairs=pairs, x0=x0, xlim=xlim)

        xy_csv_path = os.path.join(csv_dir, f"{base_file_name}_origin_xy.csv")
        pd.DataFrame({"Time (s)": x, LABELS["en"][kind]: v}).to_csv(xy_csv_path, index=False)

        out_png_name = f"{date_tag}_uet_draft_plot_{base_file_name}_v1.png"
        out_png_path = os.path.join(draft_png_dir, out_png_name)
        origin_plot([res], [raw_name], kind, lang, out_png_path, show_pp=True)

        unit = "V" if kind == "Voc" else "\u00b5A"
        if len(pairs) > 0:
            m_pp = float(pairs["pp"].mean())
            s_pp = float(pairs["pp"].std(ddof=1)) if len(pairs) > 1 else np.nan
            m_pos = float(pairs["max"].mean())
            m_neg = float(pairs["min"].mean())
            n_cyc = len(pairs)
        else:
            m_pp = float(np.ptp(v))
            s_pp = np.nan
            m_pos = float(np.max(v))
            m_neg = float(np.min(v))
            n_cyc = 0

        summary_rows.append({
            "Sample": raw_name,
            "File_Tag": base_file_name,
            "Type": kind,
            "Peak_Peak_Mean": round(m_pp, 3),
            "Peak_Peak_Std": round(s_pp, 3) if not np.isnan(s_pp) else np.nan,
            "Unit": unit,
            "N_cycles": n_cyc,
            "Peak_Pos_Mean": round(m_pos, 3),
            "Peak_Neg_Mean": round(m_neg, 3),
            "Baseline_DC": round(info["baseline"], 4),
            "Window_Start_s": round(info["tmin"], 3),
            "Window_End_s": round(info["tmax"], 3),
            "Image_File": out_png_name,
        })

    sum_df = pd.DataFrame(summary_rows)
    sum_csv_path = os.path.join(
        out_dir, f"{date_tag}_uet_dataset_nips_peak_peak_{len(summary_rows)}series_v1.csv"
    )
    sum_df.to_csv(sum_csv_path, index=False, encoding="utf-8-sig")
    return sum_df


# ==============================================================================
# 7. CLI ENTRYPOINT
# ==============================================================================
def main():
    parser = argparse.ArgumentParser(
        description="TENG Electrical Output Analyzer (Voc, Isc, Isc -> Qsc, sigma)"
    )
    subparsers = parser.add_subparsers(dest="mode", required=True)

    # Single file mode
    p_single = subparsers.add_parser("single", help="Analyze a single CSV file")
    p_single.add_argument("--input", required=True, help="Path to oscilloscope CSV file")
    p_single.add_argument("--kind", choices=["auto", "Voc", "Isc", "Qsc"], default="auto")
    p_single.add_argument("--outdir", default=".", help="Output directory")
    p_single.add_argument("--t-start", type=float, default=DEFAULT_T_START)
    p_single.add_argument("--t-end", type=float, default=DEFAULT_T_END)
    p_single.add_argument("--lang", choices=["vi", "en"], default="vi")
    p_single.add_argument("--qsc-diag", action="store_true", help="Also generate 4-panel Qsc diagnostic figure")

    # Folder batch mode
    p_folder = subparsers.add_parser("folder", help="Batch-process a folder of CSV files")
    p_folder.add_argument("--input", required=True, help="Folder path containing CSV files")
    p_folder.add_argument("--outdir", default=None, help="Output directory for PNGs")
    p_folder.add_argument("--t-start", type=float, default=DEFAULT_T_START)
    p_folder.add_argument("--t-end", type=float, default=DEFAULT_T_END)
    p_folder.add_argument("--lang", choices=["vi", "en"], default="vi")

    # Excel batch mode
    p_excel = subparsers.add_parser("excel", help="Batch-process a multi-column Excel file")
    p_excel.add_argument("--input", required=True, help="Path to Excel (.xlsx) file")
    p_excel.add_argument("--outdir", required=True, help="Output directory")
    p_excel.add_argument("--t-start", type=float, default=DEFAULT_T_START)
    p_excel.add_argument("--t-end", type=float, default=DEFAULT_T_END)
    p_excel.add_argument("--lang", choices=["vi", "en"], default="vi")

    args = parser.parse_args()
    win = (args.t_start, args.t_end)

    if args.mode == "single":
        os.makedirs(args.outdir, exist_ok=True)
        base = os.path.splitext(os.path.basename(args.input))[0]
        res = analyze_file(args.input, args.kind, window=win, preview=False)
        det_kind = res["kind"]
        out_png = os.path.join(args.outdir, f"{det_kind}_{base}_origin_plot.png")
        origin_plot([res], [base], det_kind, args.lang, out_png, show_pp=True)
        conf = res.get("confidence", "UNKNOWN")
        reason = res.get("detect_reason", "")
        print(f"Classification: {det_kind} [Confidence: {conf}] ({reason})")
        if conf in ("MODERATE", "AMBIGUOUS"):
            print("  -> Note: Classification is probabilistic. If ambiguous, please verify with experimenter or specify --kind Voc/Isc.")
        print(f"Saved Origin plot: {out_png}")
        if args.qsc_diag or det_kind in ("Isc", "Qsc"):
            cycles_df, sub_q, dom_pol = analyze_teng_qsc_cycles(res["sub"])
            diag_png = os.path.join(args.outdir, f"Qsc_{base}_4panel_diagnostic.png")
            plot_qsc_4panel(res["full_df"], sub_q, cycles_df, res["info"], base, diag_png, lang=args.lang)
            cycles_csv = os.path.join(args.outdir, f"Qsc_{base}_cycles.csv")
            cycles_df.to_csv(cycles_csv, index=False, encoding="utf-8-sig")
            print(f"Saved 4-panel Qsc diagnostic: {diag_png}")
            print(f"Saved Qsc cycle table: {cycles_csv}")

    elif args.mode == "folder":
        df = batch_process_folder(args.input, out_dir=args.outdir, window=win, lang=args.lang)
        print(df.to_string(index=False))

    elif args.mode == "excel":
        df = batch_process_excel(args.input, out_dir=args.outdir, window=win, lang=args.lang)
        print(df.to_string(index=False))


if __name__ == "__main__":
    main()

